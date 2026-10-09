#!/usr/bin/env python3
"""
Méta-moteur AAP — Scraper web en direct.
Scrape des sites publics, extrait les appels à projets via LM Studio (Qwen 2.5),
stocke en SQLite et exporte en XLSX + JSON.
"""

import argparse
import json
import logging
import re
import sqlite3
import sys
import time
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from tqdm import tqdm

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

LMSTUDIO_BASE   = "http://localhost:1234/v1"
DEFAULT_MODEL   = "qwen2.5-7b-instruct"
OUTPUT_DIR      = Path(r"C:\Users\pierr\Documents\Arcadia\AAP_moteur")
SOURCES_FILE    = Path(__file__).parent / "sources.json"

# Context LM Studio configuré à 10 000 tokens — on peut envoyer plus de texte
CHUNK_SIZE      = 6_000
REQUEST_DELAY   = 1.5   # secondes entre requêtes HTTP (politesse serveurs)
REQUEST_TIMEOUT = 25    # timeout scraping pages web

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/124.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "fr-FR,fr;q=0.9",
}

# Encodage UTF-8 forcé sur stdout Windows
_stdout_utf8 = open(sys.stdout.fileno(), mode="w", encoding="utf-8", closefd=False)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    handlers=[logging.StreamHandler(_stdout_utf8)],
)
log = logging.getLogger(__name__)

KEYWORDS_AAP = [
    "appel a projets", "appel à projets",
    "appel a candidatures", "appel à candidatures",
    "appel a manifestation", "appel à manifestation",
    "aap", "ami", "financement", "subvention",
    "depot de dossier", "dépôt de dossier",
    "candidature", "appels-a-projets",
]

EXTENSIONS_FICHIERS = {".pdf", ".docx", ".xlsx", ".odt", ".zip", ".doc", ".pptx"}

# ---------------------------------------------------------------------------
# Prompt système — optimisé pour Qwen 2.5 7B
# ---------------------------------------------------------------------------

SYSTEM_PROMPT = """\
Tu es un assistant spécialisé dans l'extraction d'informations sur les appels à projets français.
Analyse le texte fourni et extrais les informations structurées.
Réponds UNIQUEMENT avec du JSON valide, sans markdown, sans texte avant ou après.

Si le texte ne contient pas d'appel à projets : {"titre": null}

Sinon, structure ta réponse ainsi :
{
  "titre": "Titre complet de l'appel à projets",
  "type_dispositif": "Appel à projets | Appel à candidatures | AMI | Appel à manifestation d'intérêt",
  "accroche": "Résumé en 2 phrases maximum (150 caractères max)",
  "objectifs": "Objectifs en 5 points max, chacun en une ligne courte",
  "description_complete": "Description en 300 mots maximum",
  "conditions_eligibilite": "Critères d'éligibilité en 5 points max",
  "modalites_candidature": "Étapes de candidature en 3 points max",
  "porteur_nom": "Nom complet de l'organisme qui lance l'appel",
  "territoire": "Zone géographique concernée (ex: Grand Est, National, Île-de-France)",
  "type_public": "Type de structure ou de personne visée (ex: Etablissements sanitaires, Associations, PME)",
  "thematiques": ["thématique1", "thématique2"],
  "montant_txt": "Montant ou fourchette exprimée en texte (ex: jusqu'à 200 000 euros sur 3 ans)",
  "montant_min": null,
  "montant_max": null,
  "date_publication": "JJ/MM/AAAA ou null",
  "date_ouverture": "JJ/MM/AAAA ou null",
  "date_cloture": "JJ/MM/AAAA ou null",
  "url_depot": "URL complète du formulaire ou portail de dépôt en ligne, ou null",
  "contact_email": "adresse email ou null",
  "contact_tel": "numéro de téléphone ou null",
  "fichiers": [
    {
      "label": "Nom descriptif (ex: Cahier des charges, Annexe 1 - Fiche de poste ASTP)",
      "nom_fichier": "nom-du-fichier.pdf",
      "extension": ".pdf",
      "taille_txt": "4,95 Mo ou vide si inconnu",
      "type_fichier": "CAHIER_DES_CHARGES | ANNEXE | FORMULAIRE | NOTICE | REGLEMENT | AUTRE",
      "groupe": "Nom du groupe si plusieurs catégories (ex: ASTP, Universitaire) ou null",
      "url": "URL absolue du fichier ou null"
    }
  ],
  "echeances": [
    {
      "etape": "Libellé de l'étape (ex: Date limite de dépôt des candidatures)",
      "date_txt": "Date en texte tel qu'écrit sur la page",
      "date_iso": "AAAA-MM-JJ ou null",
      "heure": "14h00 ou null si non précisée"
    }
  ]
}"""


# ---------------------------------------------------------------------------
# Base de données SQLite
# ---------------------------------------------------------------------------

SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS aap (
    id                     INTEGER PRIMARY KEY AUTOINCREMENT,
    slug                   TEXT UNIQUE,
    titre                  TEXT NOT NULL,
    type_dispositif        TEXT,
    accroche               TEXT,
    objectifs              TEXT,
    description_complete   TEXT,
    conditions_eligibilite TEXT,
    modalites_candidature  TEXT,
    porteur_nom            TEXT,
    territoire             TEXT,
    type_public            TEXT,
    thematiques            TEXT,
    montant_txt            TEXT,
    montant_min            INTEGER,
    montant_max            INTEGER,
    date_publication       TEXT,
    date_ouverture         TEXT,
    date_cloture           TEXT,
    url_depot              TEXT,
    url_source             TEXT UNIQUE,
    contact_email          TEXT,
    contact_tel            TEXT,
    source_nom             TEXT,
    statut                 TEXT CHECK(statut IN ('OUVERT','CLOTURE','A_VENIR')),
    date_extraction        TEXT,
    date_maj               TEXT
);

CREATE TABLE IF NOT EXISTS fichiers (
    id           INTEGER PRIMARY KEY AUTOINCREMENT,
    aap_id       INTEGER REFERENCES aap(id) ON DELETE CASCADE,
    label        TEXT,
    nom_fichier  TEXT,
    extension    TEXT,
    taille_txt   TEXT,
    type_fichier TEXT,
    groupe       TEXT,
    url_distante TEXT,
    ordre        INTEGER
);

CREATE TABLE IF NOT EXISTS echeances (
    id       INTEGER PRIMARY KEY AUTOINCREMENT,
    aap_id   INTEGER REFERENCES aap(id) ON DELETE CASCADE,
    etape    TEXT,
    date_txt TEXT,
    date_iso TEXT,
    heure    TEXT,
    ordre    INTEGER
);

CREATE TABLE IF NOT EXISTS urls_traitees (
    url             TEXT PRIMARY KEY,
    statut          TEXT,
    date_traitement TEXT
);

CREATE VIRTUAL TABLE IF NOT EXISTS aap_fts USING fts5(
    titre, accroche, objectifs, description_complete,
    porteur_nom, territoire, thematiques, type_public,
    content='aap', content_rowid='id'
);

CREATE INDEX IF NOT EXISTS idx_statut     ON aap(statut);
CREATE INDEX IF NOT EXISTS idx_cloture    ON aap(date_cloture);
CREATE INDEX IF NOT EXISTS idx_porteur    ON aap(porteur_nom);
CREATE INDEX IF NOT EXISTS idx_territoire ON aap(territoire);
"""


def init_db(db_path: Path) -> sqlite3.Connection:
    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA foreign_keys = ON")
    conn.executescript(SCHEMA_SQL)
    conn.commit()
    return conn


def url_deja_traitee(conn: sqlite3.Connection, url: str) -> bool:
    return conn.execute(
        "SELECT 1 FROM urls_traitees WHERE url = ?", (url,)
    ).fetchone() is not None


def marquer_url(conn: sqlite3.Connection, url: str, statut: str):
    conn.execute(
        "INSERT OR REPLACE INTO urls_traitees(url, statut, date_traitement) VALUES (?,?,?)",
        (url, statut, datetime.now().isoformat()),
    )
    conn.commit()


def generer_slug(titre: str, conn: sqlite3.Connection) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", titre.lower().strip())[:80].strip("-")
    slug, i = base, 2
    while conn.execute("SELECT 1 FROM aap WHERE slug = ?", (slug,)).fetchone():
        slug = f"{base}-{i}"
        i += 1
    return slug


def normaliser_champ_texte(valeur) -> str | None:
    """Convertit une liste en texte si le modèle a renvoyé un tableau au lieu d'une chaîne."""
    if isinstance(valeur, list):
        return "\n".join(str(v) for v in valeur)
    return valeur


def inserer_aap(conn: sqlite3.Connection, data: dict,
                source_nom: str, url_source: str) -> int | None:
    slug   = generer_slug(data.get("titre", "aap"), conn)
    statut = calculer_statut(data.get("date_ouverture"), data.get("date_cloture"))
    now    = datetime.now().isoformat()
    cur = conn.execute("""
        INSERT OR IGNORE INTO aap (
            slug, titre, type_dispositif, accroche, objectifs,
            description_complete, conditions_eligibilite, modalites_candidature,
            porteur_nom, territoire, type_public, thematiques,
            montant_txt, montant_min, montant_max,
            date_publication, date_ouverture, date_cloture,
            url_depot, url_source, contact_email, contact_tel,
            source_nom, statut, date_extraction, date_maj
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    """, (
        slug,
        normaliser_champ_texte(data.get("titre")),
        normaliser_champ_texte(data.get("type_dispositif")),
        normaliser_champ_texte(data.get("accroche")),
        normaliser_champ_texte(data.get("objectifs")),
        normaliser_champ_texte(data.get("description_complete")),
        normaliser_champ_texte(data.get("conditions_eligibilite")),
        normaliser_champ_texte(data.get("modalites_candidature")),
        normaliser_champ_texte(data.get("porteur_nom")),
        normaliser_champ_texte(data.get("territoire")),
        normaliser_champ_texte(data.get("type_public")),
        json.dumps(data.get("thematiques", []), ensure_ascii=False),
        normaliser_champ_texte(data.get("montant_txt")),
        data.get("montant_min"),
        data.get("montant_max"),
        normaliser_champ_texte(data.get("date_publication")),
        normaliser_champ_texte(data.get("date_ouverture")),
        normaliser_champ_texte(data.get("date_cloture")),
        normaliser_champ_texte(data.get("url_depot")),
        url_source,
        normaliser_champ_texte(data.get("contact_email")),
        normaliser_champ_texte(data.get("contact_tel")),
        source_nom, statut, now, now,
    ))
    conn.commit()
    return cur.lastrowid if cur.lastrowid else None


def inserer_fichiers(conn: sqlite3.Connection, aap_id: int, fichiers: list):
    for i, f in enumerate(fichiers):
        conn.execute("""
            INSERT INTO fichiers
            (aap_id, label, nom_fichier, extension, taille_txt,
             type_fichier, groupe, url_distante, ordre)
            VALUES (?,?,?,?,?,?,?,?,?)
        """, (
            aap_id, f.get("label"), f.get("nom_fichier"), f.get("extension"),
            f.get("taille_txt"), f.get("type_fichier", "AUTRE"),
            f.get("groupe"), f.get("url"), i,
        ))
    conn.commit()


def inserer_echeances(conn: sqlite3.Connection, aap_id: int, echeances: list):
    for i, e in enumerate(echeances):
        conn.execute("""
            INSERT INTO echeances (aap_id, etape, date_txt, date_iso, heure, ordre)
            VALUES (?,?,?,?,?,?)
        """, (aap_id, e.get("etape"), e.get("date_txt"), e.get("date_iso"),
              e.get("heure"), i))
    conn.commit()


def rebuild_fts(conn: sqlite3.Connection):
    conn.execute("INSERT INTO aap_fts(aap_fts) VALUES('rebuild')")
    conn.commit()


# ---------------------------------------------------------------------------
# Calcul statut AAP
# ---------------------------------------------------------------------------

def calculer_statut(date_ouverture: str | None, date_cloture: str | None) -> str:
    today = date.today()

    def parse(d: str | None) -> date | None:
        if not d or str(d).lower() in ("null", "none", ""):
            return None
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(d, fmt).date()
            except ValueError:
                continue
        return None

    cloture   = parse(date_cloture)
    ouverture = parse(date_ouverture)

    if cloture and cloture < today:
        return "CLOTURE"
    if ouverture and ouverture > today:
        return "A_VENIR"
    return "OUVERT"


# ---------------------------------------------------------------------------
# Scraping web
# ---------------------------------------------------------------------------

def get_page(url: str) -> str | None:
    try:
        resp = requests.get(url, headers=HEADERS, timeout=REQUEST_TIMEOUT)
        resp.raise_for_status()
        return resp.text
    except Exception as e:
        log.warning("Erreur HTTP [%s] : %s", url, e)
        return None


def nettoyer_html(html: str) -> str:
    """Supprime les balises inutiles et retourne le texte brut nettoyé."""
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "nav", "header", "footer",
                     "noscript", "iframe", "aside", "form", "button"]):
        tag.decompose()
    texte = soup.get_text(separator="\n")
    lignes, propre, prev_vide = texte.splitlines(), [], False
    for ligne in lignes:
        ligne = ligne.strip()
        if not ligne:
            if not prev_vide:
                propre.append("")
            prev_vide = True
        else:
            propre.append(ligne)
            prev_vide = False
    return "\n".join(propre)


def extraire_fichiers_html(html: str, base_url: str) -> list[dict]:
    """Détecte les liens vers des fichiers téléchargeables dans la page."""
    soup = BeautifulSoup(html, "lxml")
    fichiers, seen = [], set()
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        url_abs = urljoin(base_url, href)
        ext = Path(urlparse(url_abs).path).suffix.lower()
        if ext not in EXTENSIONS_FICHIERS or url_abs in seen:
            continue
        seen.add(url_abs)
        label = a.get_text(strip=True) or Path(url_abs).name
        taille = ""
        if a.parent:
            m = re.search(r"\(?\s*([\d,\.]+\s*[KkMmGg]o)\s*\)?", a.parent.get_text())
            if m:
                taille = m.group(1).strip()
        fichiers.append({
            "label": label,
            "nom_fichier": Path(urlparse(url_abs).path).name,
            "extension": ext,
            "taille_txt": taille,
            "url": url_abs,
            "type_fichier": "AUTRE",
            "groupe": None,
        })
    return fichiers


def contient_keywords(texte: str) -> bool:
    lower = texte.lower()
    return any(kw in lower for kw in KEYWORDS_AAP)


def decouvrir_urls_aap(source: dict) -> list[str]:
    """Crawle la page listing d'une source et retourne les URLs d'AAP trouvées."""
    urls = list(source.get("urls_directes", []))
    url_liste = source.get("url_liste")
    if not url_liste:
        return urls

    html = get_page(url_liste)
    if not html:
        log.warning("Impossible d'accéder à la liste : %s", url_liste)
        return urls

    soup    = BeautifulSoup(html, "lxml")
    domaine = source.get("domaine", "")
    filtres = source.get("filtre_url_keywords", [])

    for a in soup.find_all("a", href=True):
        url_abs = urljoin(url_liste, a["href"])
        parsed  = urlparse(url_abs)
        if domaine and domaine not in parsed.netloc:
            continue
        if filtres and not any(kw in url_abs.lower() for kw in filtres):
            continue
        if url_abs not in urls:
            urls.append(url_abs)

    log.info("Source '%s' : %d URLs découvertes", source["nom"], len(urls))
    return urls


# ---------------------------------------------------------------------------
# LM Studio — appels API
# ---------------------------------------------------------------------------

def lister_modeles() -> list[str]:
    resp = requests.get(f"{LMSTUDIO_BASE}/models", timeout=10)
    resp.raise_for_status()
    return [m["id"] for m in resp.json().get("data", [])]


def choisir_modele(model_arg: str | None) -> str:
    if model_arg:
        return model_arg
    modeles = lister_modeles()
    if not modeles:
        log.error("Aucun modèle disponible dans LM Studio.")
        sys.exit(1)
    print("\nModèles disponibles dans LM Studio :")
    for i, m in enumerate(modeles, 1):
        print(f"  {i}. {m}")
    choix = input("Numéro du modèle à utiliser : ").strip()
    try:
        return modeles[int(choix) - 1]
    except (ValueError, IndexError):
        log.error("Choix invalide.")
        sys.exit(1)


def appeler_lmstudio(texte: str, model: str) -> dict:
    """Envoie le texte à LM Studio et retourne le JSON extrait.
    Réduit automatiquement la taille si le modèle renvoie une erreur 400."""
    tailles = [CHUNK_SIZE, CHUNK_SIZE * 2 // 3, CHUNK_SIZE // 2, CHUNK_SIZE // 4]
    derniere_erreur = None
    for taille in tailles:
        payload = {
            "model": model,
            "messages": [
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user",   "content": texte[:taille]},
            ],
            "temperature": 0.1,
            "max_tokens": 2000,
        }
        try:
            resp = requests.post(
                f"{LMSTUDIO_BASE}/chat/completions",
                json=payload,
                timeout=420,
            )
            if resp.status_code == 400:
                log.warning("Contexte trop long (%d chars) — on réduit...", taille)
                derniere_erreur = f"400 sur {taille} chars"
                continue
            resp.raise_for_status()
            contenu = resp.json()["choices"][0]["message"]["content"].strip()
            return extraire_json(contenu)
        except (requests.exceptions.Timeout, requests.exceptions.ConnectionError) as e:
            derniere_erreur = str(e)
            log.warning("Erreur réseau LM Studio : %s", e)
            break
        except json.JSONDecodeError as e:
            derniere_erreur = f"JSON invalide : {e}"
            log.warning("Réponse non-JSON de LM Studio : %s", e)
            break
    raise RuntimeError(f"LM Studio indisponible : {derniere_erreur}")


def extraire_json(contenu: str) -> dict:
    """Extrait le JSON de la réponse, tolère les troncatures dues au max_tokens."""
    # Isole le bloc JSON
    match = re.search(r"\{.*\}", contenu, re.DOTALL)
    if match:
        contenu = match.group(0)

    # Tentative directe
    try:
        return json.loads(contenu)
    except json.JSONDecodeError:
        pass

    # JSON tronqué : on tente de récupérer les champs scalaires par regex
    log.warning("JSON tronque, extraction partielle des champs...")
    data: dict = {}
    for champ in ("titre", "type_dispositif", "accroche", "porteur_nom",
                  "territoire", "type_public", "montant_txt",
                  "date_publication", "date_ouverture", "date_cloture",
                  "url_depot", "contact_email", "contact_tel"):
        m = re.search(rf'"{champ}"\s*:\s*"([^"]*)"', contenu)
        if m:
            data[champ] = m.group(1)
    # Champs tableau textuels (thematiques)
    m = re.search(r'"thematiques"\s*:\s*\[([^\]]*)\]', contenu)
    if m:
        data["thematiques"] = [t.strip().strip('"') for t in m.group(1).split(",") if t.strip()]
    # Garantit les listes manquantes
    data.setdefault("fichiers", [])
    data.setdefault("echeances", [])
    return data


# ---------------------------------------------------------------------------
# Fusion fichiers LLM + HTML
# ---------------------------------------------------------------------------

def fusionner_fichiers(fichiers_llm: list[dict], fichiers_html: list[dict]) -> list[dict]:
    """Complète les fichiers détectés par le LLM avec ceux trouvés dans le HTML."""
    noms_llm = {f.get("nom_fichier", "").lower() for f in fichiers_llm}
    for f in fichiers_html:
        if f["nom_fichier"].lower() not in noms_llm:
            fichiers_llm.append(f)
    return fichiers_llm


# ---------------------------------------------------------------------------
# Pipeline de traitement d'une URL
# ---------------------------------------------------------------------------

def traiter_url(url: str, source_nom: str, model: str,
                conn: sqlite3.Connection, log_file) -> str:
    """Scrape, extrait, insère. Retourne 'OK', 'IGNOREE' ou 'ERREUR'."""

    if url_deja_traitee(conn, url):
        log_file.write(f"{url} | IGNOREE | deja traitee\n")
        return "IGNOREE"

    html = get_page(url)
    if not html:
        log_file.write(f"{url} | ERREUR | page inaccessible\n")
        marquer_url(conn, url, "ERREUR")
        return "ERREUR"

    texte = nettoyer_html(html)
    if not contient_keywords(texte):
        log_file.write(f"{url} | IGNOREE | aucun mot-cle AAP\n")
        marquer_url(conn, url, "IGNOREE")
        return "IGNOREE"

    try:
        data = appeler_lmstudio(texte, model)
    except Exception as e:
        log.warning("LM Studio échec [%s] : %s", url, e)
        log_file.write(f"{url} | ERREUR | LM Studio : {e}\n")
        marquer_url(conn, url, "ERREUR")
        return "ERREUR"

    if not data.get("titre"):
        log_file.write(f"{url} | IGNOREE | aucun AAP detecte par le modele\n")
        marquer_url(conn, url, "IGNOREE")
        return "IGNOREE"

    fichiers_html = extraire_fichiers_html(html, url)
    data["fichiers"] = fusionner_fichiers(data.get("fichiers", []), fichiers_html)

    try:
        aap_id = inserer_aap(conn, data, source_nom, url)
        if aap_id:
            inserer_fichiers(conn, aap_id, data["fichiers"])
            inserer_echeances(conn, aap_id, data.get("echeances", []))
        log.info("OK : '%s'", data["titre"])
        log_file.write(f"{url} | OK | {data['titre']}\n")
        marquer_url(conn, url, "OK")
        return "OK"
    except Exception as e:
        log.warning("Erreur DB [%s] : %s", url, e)
        log_file.write(f"{url} | ERREUR | DB : {e}\n")
        marquer_url(conn, url, "ERREUR")
        return "ERREUR"


# ---------------------------------------------------------------------------
# Export XLSX
# ---------------------------------------------------------------------------

STATUT_COULEURS = {
    "OUVERT":  "C6EFCE",
    "CLOTURE": "FFC7CE",
    "A_VENIR": "FFEB9C",
}


def export_xlsx(conn: sqlite3.Connection, xlsx_path: Path):
    wb = Workbook()

    # Onglet Synthèse
    ws = wb.active
    ws.title = "Synthese"
    entetes = ["ID", "Titre", "Porteur", "Territoire", "Type public",
               "Date ouverture", "Date cloture", "Montant", "Statut",
               "Source", "Fichiers", "URL source"]
    ws.append(entetes)
    for cell in ws[1]:
        cell.font      = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    rows = conn.execute("""
        SELECT id, titre, porteur_nom, territoire, type_public,
               date_ouverture, date_cloture, montant_txt, statut,
               source_nom, url_source
        FROM aap ORDER BY date_cloture ASC NULLS LAST
    """).fetchall()

    for row in rows:
        aap_id = row[0]
        fichiers = conn.execute(
            "SELECT label FROM fichiers WHERE aap_id = ? ORDER BY ordre", (aap_id,)
        ).fetchall()
        fichiers_str = ", ".join(f[0] for f in fichiers if f[0])
        ws.append(list(row[:11]) + [fichiers_str])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    col_statut = get_column_letter(entetes.index("Statut") + 1)
    if ws.max_row > 1:
        for val, hex_color in STATUT_COULEURS.items():
            fill = PatternFill(start_color=hex_color, end_color=hex_color,
                               fill_type="solid")
            ws.conditional_formatting.add(
                f"{col_statut}2:{col_statut}{ws.max_row}",
                CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill),
            )

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 70)

    # Onglet Fichiers
    ws2 = wb.create_sheet("Fichiers")
    ws2.append(["ID AAP", "Titre AAP", "Label", "Nom fichier",
                "Extension", "Taille", "Type", "Groupe", "URL"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for r in conn.execute("""
        SELECT f.aap_id, a.titre, f.label, f.nom_fichier, f.extension,
               f.taille_txt, f.type_fichier, f.groupe, f.url_distante
        FROM fichiers f JOIN aap a ON a.id = f.aap_id
        ORDER BY f.aap_id, f.ordre
    """).fetchall():
        ws2.append(list(r))

    # Onglet Echeances
    ws3 = wb.create_sheet("Echeances")
    ws3.append(["ID AAP", "Titre AAP", "Etape", "Date texte", "Date ISO", "Heure"])
    for cell in ws3[1]:
        cell.font = Font(bold=True)
    for r in conn.execute("""
        SELECT e.aap_id, a.titre, e.etape, e.date_txt, e.date_iso, e.heure
        FROM echeances e JOIN aap a ON a.id = e.aap_id
        ORDER BY e.aap_id, e.ordre
    """).fetchall():
        ws3.append(list(r))

    wb.save(xlsx_path)
    log.info("XLSX exporte -> %s", xlsx_path)


# ---------------------------------------------------------------------------
# Export JSON (pour l'API du méta-moteur)
# ---------------------------------------------------------------------------

def _colonnes(conn: sqlite3.Connection, table: str) -> list[str]:
    return [d[0] for d in conn.execute(f"SELECT * FROM {table} LIMIT 0").description]


def export_json(conn: sqlite3.Connection, json_path: Path):
    cols_aap      = _colonnes(conn, "aap")
    cols_fichiers = _colonnes(conn, "fichiers")
    cols_ech      = _colonnes(conn, "echeances")
    aaps = []
    for row in conn.execute("SELECT * FROM aap ORDER BY date_cloture ASC").fetchall():
        aap = dict(zip(cols_aap, row))
        aap["thematiques"] = json.loads(aap.get("thematiques") or "[]")
        aap["fichiers"] = [
            dict(zip(cols_fichiers, f))
            for f in conn.execute(
                "SELECT * FROM fichiers WHERE aap_id = ? ORDER BY ordre", (aap["id"],)
            ).fetchall()
        ]
        aap["echeances"] = [
            dict(zip(cols_ech, e))
            for e in conn.execute(
                "SELECT * FROM echeances WHERE aap_id = ? ORDER BY ordre", (aap["id"],)
            ).fetchall()
        ]
        aaps.append(aap)

    json_path.write_text(
        json.dumps(
            {"aaps": aaps, "total": len(aaps), "export_date": datetime.now().isoformat()},
            ensure_ascii=False, indent=2,
        ),
        encoding="utf-8",
    )
    log.info("JSON exporte -> %s (%d AAP)", json_path, len(aaps))


# ---------------------------------------------------------------------------
# CLI et orchestration
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Scraper AAP web - meta-moteur de recherche",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Exemples :\n"
            f'  python aap_web_scraper.py --model "{DEFAULT_MODEL}"\n'
            f'  python aap_web_scraper.py --url "https://exemple.fr/aap" --model "{DEFAULT_MODEL}"\n'
            f'  python aap_web_scraper.py --sources sources.json --model "{DEFAULT_MODEL}"'
        ),
    )
    p.add_argument("--model",   default=None,          help=f"ID modèle LM Studio (défaut auto-sélection, recommandé: {DEFAULT_MODEL})")
    p.add_argument("--url",     default=None,          help="Scraper une seule URL")
    p.add_argument("--sources", default=str(SOURCES_FILE), help="Fichier sources.json (défaut: sources.json)")
    p.add_argument("--reset",   action="store_true",   help="Réinitialiser l'historique des URLs traitées")
    return p.parse_args()


def main():
    args = parse_args()

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    db_path   = OUTPUT_DIR / "aap_moteur.db"
    xlsx_path = OUTPUT_DIR / "aap_moteur.xlsx"
    json_path = OUTPUT_DIR / "aap_moteur.json"
    log_path  = OUTPUT_DIR / "scraping.log"

    model = choisir_modele(args.model)
    log.info("Modele selectionne : %s", model)

    conn = init_db(db_path)

    if args.reset:
        conn.execute("DELETE FROM urls_traitees")
        conn.commit()
        log.info("Historique des URLs reinitialise.")

    if args.url:
        sources = [{"nom": "URL directe", "urls_directes": [args.url],
                    "url_liste": None, "filtre_url_keywords": [], "domaine": ""}]
    else:
        sources_path = Path(args.sources)
        if not sources_path.exists():
            log.error("Fichier sources introuvable : %s", sources_path)
            sys.exit(1)
        sources = json.loads(sources_path.read_text(encoding="utf-8"))

    toutes_urls: list[tuple[str, str]] = []
    for source in sources:
        for url in decouvrir_urls_aap(source):
            toutes_urls.append((url, source["nom"]))

    log.info("Total URLs a traiter : %d", len(toutes_urls))
    if not toutes_urls:
        log.warning("Aucune URL a traiter. Verifiez votre sources.json ou utilisez --url.")
        return

    compteurs: dict[str, int] = {"OK": 0, "IGNOREE": 0, "ERREUR": 0}

    with open(log_path, "a", encoding="utf-8") as lf:
        for url, source_nom in tqdm(toutes_urls, desc="Scraping AAP", unit="page"):
            statut = traiter_url(url, source_nom, model, conn, lf)
            compteurs[statut] = compteurs.get(statut, 0) + 1
            time.sleep(REQUEST_DELAY)

    rebuild_fts(conn)

    log.info("--- Resultats ---")
    log.info("  Extraits  : %d", compteurs["OK"])
    log.info("  Ignores   : %d", compteurs["IGNOREE"])
    log.info("  Erreurs   : %d", compteurs["ERREUR"])

    export_xlsx(conn, xlsx_path)
    export_json(conn, json_path)
    conn.close()

    log.info("Tous les fichiers sont dans : %s", OUTPUT_DIR)


if __name__ == "__main__":
    main()
