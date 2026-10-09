#!/usr/bin/env python3
"""
Extracteur automatique d'appels à projets depuis une archive HTTrack.
Analyse les fichiers HTML locaux, extrait les AAP via LM Studio (Qwen 2.5),
stocke en SQLite et exporte en XLSX.
"""

import argparse
import json
import logging
import re
import sqlite3
import sys
from datetime import date, datetime
from pathlib import Path

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from tqdm import tqdm

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

LMSTUDIO_BASE   = "http://localhost:1234/v1"
DEFAULT_MODEL   = "qwen2.5-7b-instruct"
OUTPUT_DIR      = Path(r"C:\Users\pierr\Documents\Arcadia\HTTrack\_extraction")

CHUNK_SIZE      = 6_000   # chars — adapté au contexte 10 000 tokens de Qwen
MAX_TOKENS      = 2_000   # tokens de réponse max
LM_TIMEOUT      = 420     # secondes

HTML_EXTS            = {".html", ".htm"}
DOWNLOADABLE_EXTS    = {".pdf", ".docx", ".xlsx", ".zip", ".odt", ".doc", ".pptx"}
HTTRACK_IGNORES      = {"hts-cache", "hts-log", "hts-err", "hts-ioinfo"}

KEYWORDS_AAP = [
    "appel a projets", "appel à projets",
    "appel a candidatures", "appel à candidatures",
    "appel a manifestation", "appel à manifestation",
    "aap", "ami", "financement", "subvention",
    "depot de dossier", "dépôt de dossier", "candidature",
]

# Encodage UTF-8 forcé sur stdout Windows
_stdout_utf8 = open(sys.stdout.fileno(), mode="w", encoding="utf-8", closefd=False)
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
    handlers=[logging.StreamHandler(_stdout_utf8)],
)
log = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Prompt système — Qwen 2.5 7B, contexte 10 000 tokens
# ---------------------------------------------------------------------------

SYSTEM_PROMPT = """\
Tu es un assistant spécialisé dans l'extraction d'informations sur les appels à projets français.
Analyse le texte fourni et extrais les informations structurées.
Réponds UNIQUEMENT avec du JSON valide, sans markdown, sans texte avant ou après.

Si le texte ne contient pas d'appel à projets : {"appels_a_projets": []}

Sinon :
{
  "appels_a_projets": [
    {
      "nom": "Titre complet de l'appel à projets",
      "type_dispositif": "Appel à projets | Appel à candidatures | AMI | Appel à manifestation d'intérêt",
      "accroche": "Résumé en 2 phrases maximum (150 caractères max)",
      "objectifs": "Objectifs en 5 points max, chacun en une ligne courte",
      "description_complete": "Description en 200 mots maximum",
      "criteres_eligibilite": ["critère 1", "critère 2"],
      "modalites_candidature": "Etapes de candidature en 3 points max",
      "porteur": "Nom de l'organisme porteur",
      "territoire": "Zone géographique (ex: Grand Est, National)",
      "type_public": "Type de structure visée (ex: Etablissements sanitaires)",
      "thematiques": ["thématique1", "thématique2"],
      "montant_aide": "Montant ou fourchette en texte (ex: jusqu'à 200 000 euros)",
      "montant_min": null,
      "montant_max": null,
      "date_ouverture": "JJ/MM/AAAA ou null",
      "date_cloture": "JJ/MM/AAAA ou null",
      "date_publication": "JJ/MM/AAAA ou null",
      "url_depot": "URL du formulaire de dépôt en ligne ou null",
      "contact_email": "email ou null",
      "contact_tel": "téléphone ou null",
      "fichiers": [
        {
          "label": "Nom descriptif (ex: Cahier des charges, Annexe 1 - Fiche de poste)",
          "nom_fichier": "nom-du-fichier.pdf",
          "extension": ".pdf",
          "groupe": "Groupe si plusieurs catégories (ex: ASTP) ou null"
        }
      ],
      "echeances": [
        {
          "etape": "Libellé de l'étape (ex: Date limite de dépôt)",
          "date_txt": "Date en texte tel qu'écrit sur la page",
          "date_iso": "AAAA-MM-JJ ou null",
          "heure": "14h00 ou null"
        }
      ],
      "page_source": "chemin/relatif/fichier.html"
    }
  ]
}"""

# ---------------------------------------------------------------------------
# Schéma SQLite
# ---------------------------------------------------------------------------

SCHEMA_SQL = """
CREATE TABLE IF NOT EXISTS aap (
    id                     INTEGER PRIMARY KEY AUTOINCREMENT,
    nom                    TEXT,
    type_dispositif        TEXT,
    accroche               TEXT,
    objectifs              TEXT,
    description_complete   TEXT,
    criteres_eligibilite   TEXT,
    modalites_candidature  TEXT,
    porteur                TEXT,
    territoire             TEXT,
    type_public            TEXT,
    thematiques            TEXT,
    montant_aide           TEXT,
    montant_min            INTEGER,
    montant_max            INTEGER,
    date_ouverture         TEXT,
    date_cloture           TEXT,
    date_publication       TEXT,
    url_depot              TEXT,
    contact_email          TEXT,
    contact_tel            TEXT,
    page_source            TEXT,
    statut                 TEXT CHECK(statut IN ('OUVERT','CLOTURE','A_VENIR')),
    date_extraction        TEXT
);

CREATE TABLE IF NOT EXISTS fichiers (
    id            INTEGER PRIMARY KEY AUTOINCREMENT,
    aap_id        INTEGER REFERENCES aap(id) ON DELETE CASCADE,
    label         TEXT,
    nom_fichier   TEXT,
    chemin_absolu TEXT,
    extension     TEXT,
    groupe        TEXT,
    ordre         INTEGER
);

CREATE TABLE IF NOT EXISTS echeances (
    id       INTEGER PRIMARY KEY AUTOINCREMENT,
    aap_id   INTEGER REFERENCES aap(id) ON DELETE CASCADE,
    etape    TEXT,
    date_txt TEXT,
    date_iso TEXT,
    heure    TEXT,
    ordre    INTEGER
);

CREATE TABLE IF NOT EXISTS pages_traitees (
    chemin          TEXT PRIMARY KEY,
    statut          TEXT,
    nb_aap          INTEGER,
    date_traitement TEXT
);

CREATE INDEX IF NOT EXISTS idx_aap_nom       ON aap(nom);
CREATE INDEX IF NOT EXISTS idx_aap_cloture   ON aap(date_cloture);
CREATE INDEX IF NOT EXISTS idx_aap_porteur   ON aap(porteur);
CREATE INDEX IF NOT EXISTS idx_aap_statut    ON aap(statut);
"""

# ---------------------------------------------------------------------------
# Base de données
# ---------------------------------------------------------------------------

def init_db(db_path: Path) -> sqlite3.Connection:
    conn = sqlite3.connect(db_path)
    conn.execute("PRAGMA foreign_keys = ON")
    conn.executescript(SCHEMA_SQL)
    conn.commit()
    return conn


def page_deja_traitee(conn: sqlite3.Connection, chemin: str) -> bool:
    return conn.execute(
        "SELECT 1 FROM pages_traitees WHERE chemin = ?", (chemin,)
    ).fetchone() is not None


def marquer_page(conn: sqlite3.Connection, chemin: str, statut: str, nb_aap: int):
    conn.execute(
        "INSERT OR REPLACE INTO pages_traitees(chemin, statut, nb_aap, date_traitement) "
        "VALUES (?,?,?,?)",
        (chemin, statut, nb_aap, datetime.now().isoformat()),
    )
    conn.commit()


def inserer_aap(conn: sqlite3.Connection, aap: dict) -> int:
    now = datetime.now().isoformat()
    statut = calculer_statut(aap.get("date_ouverture"), aap.get("date_cloture"))
    cur = conn.execute("""
        INSERT INTO aap (
            nom, type_dispositif, accroche, objectifs, description_complete,
            criteres_eligibilite, modalites_candidature, porteur, territoire,
            type_public, thematiques, montant_aide, montant_min, montant_max,
            date_ouverture, date_cloture, date_publication,
            url_depot, contact_email, contact_tel,
            page_source, statut, date_extraction
        ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
    """, (
        normaliser_texte(aap.get("nom")),
        normaliser_texte(aap.get("type_dispositif")),
        normaliser_texte(aap.get("accroche")),
        normaliser_texte(aap.get("objectifs")),
        normaliser_texte(aap.get("description_complete")),
        json.dumps(aap.get("criteres_eligibilite", []), ensure_ascii=False),
        normaliser_texte(aap.get("modalites_candidature")),
        normaliser_texte(aap.get("porteur")),
        normaliser_texte(aap.get("territoire")),
        normaliser_texte(aap.get("type_public")),
        json.dumps(aap.get("thematiques", []), ensure_ascii=False),
        normaliser_texte(aap.get("montant_aide")),
        aap.get("montant_min"),
        aap.get("montant_max"),
        normaliser_texte(aap.get("date_ouverture")),
        normaliser_texte(aap.get("date_cloture")),
        normaliser_texte(aap.get("date_publication")),
        normaliser_texte(aap.get("url_depot")),
        normaliser_texte(aap.get("contact_email")),
        normaliser_texte(aap.get("contact_tel")),
        aap.get("page_source"),
        statut,
        now,
    ))
    conn.commit()
    return cur.lastrowid


def inserer_fichiers(conn: sqlite3.Connection, aap_id: int, fichiers: list):
    for i, f in enumerate(fichiers):
        conn.execute("""
            INSERT INTO fichiers (aap_id, label, nom_fichier, chemin_absolu, extension, groupe, ordre)
            VALUES (?,?,?,?,?,?,?)
        """, (
            aap_id,
            f.get("label"),
            f.get("nom_fichier"),
            f.get("chemin_absolu"),
            f.get("extension"),
            f.get("groupe"),
            i,
        ))
    conn.commit()


def inserer_echeances(conn: sqlite3.Connection, aap_id: int, echeances: list):
    for i, e in enumerate(echeances):
        conn.execute("""
            INSERT INTO echeances (aap_id, etape, date_txt, date_iso, heure, ordre)
            VALUES (?,?,?,?,?,?)
        """, (aap_id, e.get("etape"), e.get("date_txt"),
              e.get("date_iso"), e.get("heure"), i))
    conn.commit()

# ---------------------------------------------------------------------------
# Utilitaires
# ---------------------------------------------------------------------------

def normaliser_texte(valeur) -> str | None:
    """Convertit une liste en texte si le modèle a renvoyé un tableau."""
    if isinstance(valeur, list):
        return "\n".join(str(v) for v in valeur)
    return valeur


def calculer_statut(date_ouverture: str | None, date_cloture: str | None) -> str:
    today = date.today()

    def parse(d: str | None) -> date | None:
        if not d or str(d).lower() in ("null", "none", ""):
            return None
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(d, fmt).date()
            except ValueError:
                continue
        return None

    cloture   = parse(date_cloture)
    ouverture = parse(date_ouverture)
    if cloture and cloture < today:
        return "CLOTURE"
    if ouverture and ouverture > today:
        return "A_VENIR"
    return "OUVERT"

# ---------------------------------------------------------------------------
# Scan du dossier HTTrack
# ---------------------------------------------------------------------------

def est_technique_httrack(path: Path) -> bool:
    for part in path.parts:
        if part.lower() in HTTRACK_IGNORES:
            return True
    name = path.name.lower()
    return name.startswith("hts-")


def trouver_fichiers_html(root: Path) -> list[Path]:
    return [
        p for p in root.rglob("*")
        if p.suffix.lower() in HTML_EXTS and not est_technique_httrack(p)
    ]


def indexer_fichiers_telechargeables(root: Path) -> dict[str, Path]:
    """Retourne {nom_fichier: chemin_absolu} pour tous les fichiers téléchargeables."""
    index = {}
    for p in root.rglob("*"):
        if p.suffix.lower() in DOWNLOADABLE_EXTS:
            index[p.name] = p
    return index

# ---------------------------------------------------------------------------
# Extraction HTML
# ---------------------------------------------------------------------------

def nettoyer_html(html_path: Path) -> str:
    """Extrait le texte brut nettoyé d'un fichier HTML local."""
    try:
        soup = BeautifulSoup(html_path.read_bytes(), "lxml")
    except Exception:
        soup = BeautifulSoup(html_path.read_text(errors="replace"), "html.parser")

    for tag in soup(["script", "style", "nav", "header", "footer",
                     "noscript", "iframe", "aside", "button"]):
        tag.decompose()

    texte = soup.get_text(separator="\n")
    lignes, propre, prev_vide = texte.splitlines(), [], False
    for ligne in lignes:
        ligne = ligne.strip()
        if not ligne:
            if not prev_vide:
                propre.append("")
            prev_vide = True
        else:
            propre.append(ligne)
            prev_vide = False
    return "\n".join(propre)


def extraire_fichiers_html(html_path: Path, index: dict[str, Path]) -> list[dict]:
    """Détecte les liens vers fichiers téléchargeables dans le HTML local
    et les croise avec l'index pour obtenir le chemin absolu."""
    try:
        soup = BeautifulSoup(html_path.read_bytes(), "lxml")
    except Exception:
        soup = BeautifulSoup(html_path.read_text(errors="replace"), "html.parser")

    fichiers, seen = [], set()
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        nom  = Path(href.split("?")[0]).name
        ext  = Path(nom).suffix.lower()
        if ext not in DOWNLOADABLE_EXTS or nom in seen:
            continue
        seen.add(nom)
        label = a.get_text(strip=True) or nom
        taille = ""
        if a.parent:
            m = re.search(r"\(?\s*([\d,\.]+\s*[KkMmGg]o)\s*\)?", a.parent.get_text())
            if m:
                taille = m.group(1).strip()
        fichiers.append({
            "label":         label,
            "nom_fichier":   nom,
            "chemin_absolu": str(index[nom]) if nom in index else None,
            "extension":     ext,
            "taille_txt":    taille,
            "groupe":        None,
        })
    return fichiers


def contient_keywords(texte: str) -> bool:
    lower = texte.lower()
    return any(kw in lower for kw in KEYWORDS_AAP)


def fusionner_fichiers(fichiers_llm: list[dict], fichiers_html: list[dict],
                       index: dict[str, Path]) -> list[dict]:
    """Fusionne : priorité aux données LLM, complète avec HTML, enrichit les chemins."""
    # Enrichit les fichiers LLM avec le chemin absolu depuis l'index
    for f in fichiers_llm:
        nom = f.get("nom_fichier", "")
        if nom in index and not f.get("chemin_absolu"):
            f["chemin_absolu"] = str(index[nom])

    noms_llm = {f.get("nom_fichier", "").lower() for f in fichiers_llm}
    for f in fichiers_html:
        if f["nom_fichier"].lower() not in noms_llm:
            fichiers_llm.append(f)
    return fichiers_llm

# ---------------------------------------------------------------------------
# LM Studio
# ---------------------------------------------------------------------------

def lister_modeles() -> list[str]:
    resp = requests.get(f"{LMSTUDIO_BASE}/models", timeout=10)
    resp.raise_for_status()
    return [m["id"] for m in resp.json().get("data", [])]


def choisir_modele(model_arg: str | None) -> str:
    if model_arg:
        return model_arg
    modeles = lister_modeles()
    if not modeles:
        log.error("Aucun modèle disponible dans LM Studio.")
        sys.exit(1)
    print("\nModèles disponibles :")
    for i, m in enumerate(modeles, 1):
        print(f"  {i}. {m}")
    choix = input("Numéro du modèle : ").strip()
    try:
        return modeles[int(choix) - 1]
    except (ValueError, IndexError):
        log.error("Choix invalide.")
        sys.exit(1)


def appeler_lmstudio(texte: str, model: str) -> dict:
    """Envoie le texte à LM Studio avec réduction automatique si contexte trop long."""
    tailles = [CHUNK_SIZE, CHUNK_SIZE * 2 // 3, CHUNK_SIZE // 2]
    derniere_erreur = None
    for taille in tailles:
        payload = {
            "model": model,
            "messages": [
                {"role": "system", "content": SYSTEM_PROMPT},
                {"role": "user",   "content": texte[:taille]},
            ],
            "temperature": 0.1,
            "max_tokens":  MAX_TOKENS,
        }
        try:
            resp = requests.post(
                f"{LMSTUDIO_BASE}/chat/completions",
                json=payload,
                timeout=LM_TIMEOUT,
            )
            if resp.status_code == 400:
                log.warning("Contexte trop long (%d chars) — on reduit...", taille)
                derniere_erreur = f"400 sur {taille} chars"
                continue
            resp.raise_for_status()
            contenu = resp.json()["choices"][0]["message"]["content"].strip()
            return extraire_json(contenu)
        except (requests.exceptions.Timeout, requests.exceptions.ConnectionError) as e:
            derniere_erreur = str(e)
            log.warning("Erreur reseau LM Studio : %s", e)
            break
        except json.JSONDecodeError as e:
            derniere_erreur = f"JSON invalide : {e}"
            log.warning("Reponse non-JSON de LM Studio : %s", e)
            break
    raise RuntimeError(f"LM Studio indisponible : {derniere_erreur}")


def extraire_json(contenu: str) -> dict:
    """Extrait le JSON de la réponse, tolère les troncatures."""
    match = re.search(r"\{.*\}", contenu, re.DOTALL)
    if match:
        contenu = match.group(0)
    try:
        return json.loads(contenu)
    except json.JSONDecodeError:
        pass

    # Fallback : extraction partielle par regex sur JSON tronqué
    log.warning("JSON tronque — extraction partielle...")
    data: dict = {"appels_a_projets": []}
    aap: dict  = {}
    for champ in ("nom", "type_dispositif", "accroche", "porteur", "territoire",
                  "type_public", "montant_aide", "date_publication",
                  "date_ouverture", "date_cloture", "url_depot",
                  "contact_email", "contact_tel", "page_source"):
        m = re.search(rf'"{champ}"\s*:\s*"([^"]*)"', contenu)
        if m:
            aap[champ] = m.group(1)
    m = re.search(r'"thematiques"\s*:\s*\[([^\]]*)\]', contenu)
    if m:
        aap["thematiques"] = [t.strip().strip('"') for t in m.group(1).split(",") if t.strip()]
    aap.setdefault("criteres_eligibilite", [])
    aap.setdefault("fichiers", [])
    aap.setdefault("echeances", [])
    if aap.get("nom"):
        data["appels_a_projets"].append(aap)
    return data

# ---------------------------------------------------------------------------
# Pipeline d'analyse d'une page
# ---------------------------------------------------------------------------

def analyser_page(html_path: Path, root: Path, index: dict[str, Path],
                  model: str) -> list[dict]:
    """Analyse une page HTML, retourne la liste des AAP trouvés."""
    texte = nettoyer_html(html_path)
    if not contient_keywords(texte):
        return []

    rel_path = str(html_path.relative_to(root))

    try:
        result = appeler_lmstudio(texte, model)
    except Exception as e:
        log.warning("LM Studio echec [%s] : %s", html_path.name, e)
        raise

    aaps = result.get("appels_a_projets", [])

    # Fichiers détectés dans le HTML (lien <a href>)
    fichiers_html = extraire_fichiers_html(html_path, index)

    for aap in aaps:
        aap["page_source"] = rel_path
        aap["fichiers"] = fusionner_fichiers(
            aap.get("fichiers", []), fichiers_html, index
        )

    return aaps


def dedupliquer(aaps: list[dict]) -> list[dict]:
    seen, uniques = set(), []
    for aap in aaps:
        cle = (aap.get("nom") or "").strip().lower()
        if cle and cle not in seen:
            seen.add(cle)
            uniques.append(aap)
    return uniques

# ---------------------------------------------------------------------------
# Export XLSX
# ---------------------------------------------------------------------------

STATUT_COULEURS = {
    "OUVERT":  "C6EFCE",
    "CLOTURE": "FFC7CE",
    "A_VENIR": "FFEB9C",
}


def export_xlsx(conn: sqlite3.Connection, xlsx_path: Path):
    wb = Workbook()

    # --- Onglet Synthèse ---
    ws = wb.active
    ws.title = "Synthese"
    entetes = ["ID", "Nom", "Porteur", "Territoire", "Type public",
               "Date ouverture", "Date cloture", "Montant", "Statut",
               "Fichiers joints", "Page source"]
    ws.append(entetes)
    for cell in ws[1]:
        cell.font      = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    rows = conn.execute("""
        SELECT id, nom, porteur, territoire, type_public,
               date_ouverture, date_cloture, montant_aide, statut, page_source
        FROM aap ORDER BY date_cloture ASC NULLS LAST
    """).fetchall()

    for row in rows:
        aap_id = row[0]
        fichiers = conn.execute(
            "SELECT COALESCE(label, nom_fichier) FROM fichiers WHERE aap_id=? ORDER BY ordre",
            (aap_id,)
        ).fetchall()
        fichiers_str = ", ".join(f[0] for f in fichiers if f[0])
        ws.append(list(row) + [fichiers_str])

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    col_statut = get_column_letter(entetes.index("Statut") + 1)
    if ws.max_row > 1:
        for val, hex_color in STATUT_COULEURS.items():
            fill = PatternFill(start_color=hex_color, end_color=hex_color,
                               fill_type="solid")
            ws.conditional_formatting.add(
                f"{col_statut}2:{col_statut}{ws.max_row}",
                CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill),
            )

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=8)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 70)

    # --- Onglet Fichiers ---
    ws2 = wb.create_sheet("Fichiers")
    ws2.append(["ID AAP", "Nom AAP", "Label", "Nom fichier",
                "Extension", "Groupe", "Chemin absolu"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for r in conn.execute("""
        SELECT f.aap_id, a.nom, f.label, f.nom_fichier,
               f.extension, f.groupe, f.chemin_absolu
        FROM fichiers f JOIN aap a ON a.id = f.aap_id
        ORDER BY f.aap_id, f.ordre
    """).fetchall():
        ws2.append(list(r))

    # --- Onglet Echeances ---
    ws3 = wb.create_sheet("Echeances")
    ws3.append(["ID AAP", "Nom AAP", "Etape", "Date texte", "Date ISO", "Heure"])
    for cell in ws3[1]:
        cell.font = Font(bold=True)
    for r in conn.execute("""
        SELECT e.aap_id, a.nom, e.etape, e.date_txt, e.date_iso, e.heure
        FROM echeances e JOIN aap a ON a.id = e.aap_id
        ORDER BY e.aap_id, e.ordre
    """).fetchall():
        ws3.append(list(r))

    wb.save(xlsx_path)
    log.info("XLSX exporte -> %s", xlsx_path)

# ---------------------------------------------------------------------------
# CLI et orchestration
# ---------------------------------------------------------------------------

def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description="Extracteur AAP depuis archive HTTrack",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=(
            "Exemples :\n"
            '  python httrack_aap_extractor.py --dossier "C:\\Users\\pierr\\Documents\\Arcadia\\HTTrack"\n'
            f'  python httrack_aap_extractor.py --dossier "..." --model "{DEFAULT_MODEL}"\n'
            '  python httrack_aap_extractor.py --dossier "..." --model "..." --reset'
        ),
    )
    p.add_argument("--dossier", required=True, help="Chemin vers l'archive HTTrack")
    p.add_argument("--model",   default=None,  help=f"ID modèle LM Studio (recommandé: {DEFAULT_MODEL})")
    p.add_argument("--reset",   action="store_true", help="Réinitialiser l'historique des pages traitées")
    return p.parse_args()


def main():
    args  = parse_args()
    root  = Path(args.dossier).resolve()
    if not root.exists():
        log.error("Dossier introuvable : %s", root)
        sys.exit(1)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    db_path   = OUTPUT_DIR / "appels_a_projets.db"
    xlsx_path = OUTPUT_DIR / "appels_a_projets.xlsx"
    log_path  = OUTPUT_DIR / "pages_analysees.log"

    model = choisir_modele(args.model)
    log.info("Modele : %s", model)

    if args.reset and db_path.exists():
        try:
            db_path.unlink()
            log.info("Base de donnees supprimee et recree (--reset).")
        except PermissionError:
            log.error("Impossible de supprimer la DB : fermez DB Browser for SQLite ou tout autre programme qui l'utilise, puis relancez.")
            sys.exit(1)

    conn = init_db(db_path)

    log.info("Scan du dossier HTTrack...")
    html_files = trouver_fichiers_html(root)
    index      = indexer_fichiers_telechargeables(root)
    log.info("%d fichiers HTML | %d fichiers telechargeable indexes", len(html_files), len(index))

    tous_aaps: list[dict] = []

    with open(log_path, "a", encoding="utf-8") as lf:
        for html_path in tqdm(html_files, desc="Analyse des pages", unit="page"):
            chemin_str = str(html_path)

            if page_deja_traitee(conn, chemin_str):
                lf.write(f"{chemin_str} | IGNOREE | deja traitee\n")
                continue

            try:
                aaps = analyser_page(html_path, root, index, model)
            except Exception as e:
                lf.write(f"{chemin_str} | ERREUR | {e}\n")
                marquer_page(conn, chemin_str, "ERREUR", 0)
                continue

            nb = len(aaps)
            if nb == 0:
                lf.write(f"{chemin_str} | IGNOREE | 0 AAP\n")
            else:
                lf.write(f"{chemin_str} | OK | {nb} AAP\n")
                tous_aaps.extend(aaps)

            marquer_page(conn, chemin_str, "OK" if nb > 0 else "IGNOREE", nb)

    log.info("%d AAP collectes — deduplication...", len(tous_aaps))
    tous_aaps = dedupliquer(tous_aaps)
    log.info("%d AAP uniques — insertion en base...", len(tous_aaps))

    for aap in tous_aaps:
        aap_id = inserer_aap(conn, aap)
        inserer_fichiers(conn, aap_id, aap.get("fichiers", []))
        inserer_echeances(conn, aap_id, aap.get("echeances", []))

    export_xlsx(conn, xlsx_path)
    conn.close()

    log.info("Termine. Resultats dans : %s", OUTPUT_DIR)
    log.info("  DB   -> %s", db_path)
    log.info("  XLSX -> %s", xlsx_path)
    log.info("  LOG  -> %s", log_path)


if __name__ == "__main__":
    main()
