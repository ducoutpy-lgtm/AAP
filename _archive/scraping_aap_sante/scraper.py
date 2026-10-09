"""
Scraper d'Appels à Projets Santé & Médico-Social
-------------------------------------------------
- Lit les URLs depuis urls.txt (modifiable librement)
- Scrape chaque site avec crawl4ai (gère le JavaScript)
- Extrait les infos via LM Studio (IA locale, gratuit)
- Stocke dans une base SQLite + exporte en Excel
- Télécharge les PDFs trouvés sur les pages
"""

import asyncio
import sqlite3
import json
import os
import re
import requests
from datetime import datetime
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urljoin, urlparse

from crawl4ai import AsyncWebCrawler
from openai import OpenAI
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ─── Configuration ────────────────────────────────────────────────────────────

URLS_FILE    = "urls.txt"       # Fichier des URLs à scraper (modifiable)
DB_PATH      = "aap_sante.db"  # Base de données SQLite
PDFS_DIR     = Path("pdfs")    # Dossier de téléchargement des PDFs
EXCEL_PATH   = "aap_sante.xlsx"

LM_STUDIO_URL = "http://localhost:1234/v1"  # Adresse de LM Studio (ne pas changer)

# ─── Chargement des URLs ──────────────────────────────────────────────────────

def load_urls(filepath):
    """Charge les URLs depuis le fichier, ignore les commentaires et lignes vides."""
    if not Path(filepath).exists():
        print(f"❌ Fichier {filepath} introuvable.")
        print(f"   Créez un fichier {filepath} avec une URL par ligne.")
        return []
    urls = []
    with open(filepath, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if line and not line.startswith("#"):
                urls.append(line)
    return urls

# ─── Base de données ──────────────────────────────────────────────────────────

def init_db():
    """Crée la base de données et les tables si elles n'existent pas."""
    conn = sqlite3.connect(DB_PATH)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS aap (
            id                   INTEGER PRIMARY KEY AUTOINCREMENT,
            titre                TEXT NOT NULL,
            description          TEXT,
            date_debut           TEXT,
            date_fin             TEXT,
            criteres_eligibilite TEXT,
            thematique           TEXT,
            url_source           TEXT NOT NULL,
            date_ajout           TEXT NOT NULL,
            fichiers             TEXT DEFAULT '[]',
            UNIQUE(titre, url_source)
        )
    """)
    conn.commit()
    conn.close()

def save_to_db(aap_list, url_source, fichiers):
    """Sauvegarde les AAP en base, ignore les doublons (même titre + même URL)."""
    conn = sqlite3.connect(DB_PATH)
    new_count = 0
    now = datetime.now().strftime("%Y-%m-%d %H:%M")

    for aap in aap_list:
        titre = (aap.get("titre") or "").strip()
        if not titre or len(titre) < 5:
            continue
        try:
            conn.execute("""
                INSERT OR IGNORE INTO aap
                (titre, description, date_debut, date_fin,
                 criteres_eligibilite, thematique, url_source, date_ajout, fichiers)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """, (
                titre,
                aap.get("description", ""),
                aap.get("date_debut", ""),
                aap.get("date_fin", ""),
                aap.get("criteres_eligibilite", ""),
                aap.get("thematique", ""),
                url_source,
                now,
                json.dumps(fichiers, ensure_ascii=False),
            ))
            if conn.execute("SELECT changes()").fetchone()[0] > 0:
                new_count += 1
        except Exception as e:
            print(f"  ⚠️  Erreur base de données : {e}")

    conn.commit()
    conn.close()
    return new_count

# ─── Extraction par LM Studio ─────────────────────────────────────────────────

def extract_with_lm_studio(markdown_text, url):
    """Envoie le texte de la page à LM Studio pour extraction structurée."""
    try:
        client = OpenAI(base_url=LM_STUDIO_URL, api_key="lm-studio")

        prompt = f"""Tu es un extracteur d'information spécialisé dans les appels à projets français.
Analyse ce texte d'une page web et extrais TOUS les appels à projets présents.

URL source : {url}

TEXTE DE LA PAGE :
{markdown_text[:7000]}

Pour chaque appel à projet trouvé, crée un objet JSON avec exactement ces champs :
- "titre"                : nom complet de l'appel à projet
- "description"          : objectif ou résumé court (2-3 phrases)
- "date_debut"           : date d'ouverture (format JJ/MM/AAAA, ou "" si inconnue)
- "date_fin"             : date limite de dépôt (format JJ/MM/AAAA, ou "" si inconnue)
- "criteres_eligibilite" : qui peut candidater (résumé court, ou "" si inconnu)
- "thematique"           : domaine concerné (ex: cancer, autonomie, prévention, handicap...)

Réponds UNIQUEMENT avec un tableau JSON valide. Aucun texte avant ou après.
Si aucun appel à projet n'est trouvé : réponds []

Exemple de réponse attendue :
[
  {{
    "titre": "AAP Innovation en santé 2025",
    "description": "Soutenir les projets innovants en santé numérique.",
    "date_debut": "01/01/2025",
    "date_fin": "30/06/2025",
    "criteres_eligibilite": "Établissements de santé et structures médico-sociales",
    "thematique": "santé numérique"
  }}
]"""

        response = client.chat.completions.create(
            model="local-model",
            messages=[
                {"role": "system", "content": "You are a JSON extractor. Output ONLY a valid JSON array, no explanation, no reasoning, no markdown."},
                {"role": "user", "content": prompt},
            ],
            temperature=0.1,
            max_tokens=4000,
        )

        content = response.choices[0].message.content.strip()
        # Nettoyer les balises markdown si présentes
        content = re.sub(r"```json\s*", "", content)
        content = re.sub(r"```\s*", "", content)
        content = content.strip()

        # Extraire le tableau JSON même si le modèle ajoute du texte autour
        # Cherche le premier '[' et le dernier ']' pour gérer le JSON complet
        start = content.find("[")
        end   = content.rfind("]")
        if start != -1 and end != -1 and end > start:
            return json.loads(content[start:end + 1])

        return json.loads(content)

    except ConnectionRefusedError:
        print("  ❌ LM Studio inaccessible. Vérifiez que le serveur est démarré sur le port 1234.")
        return []
    except json.JSONDecodeError as e:
        print(f"  ⚠️  Réponse LM Studio non parseable : {e}")
        print(f"  ↳  Début de la réponse : {content[:200]!r}")
        return []
    except Exception as e:
        print(f"  ⚠️  Erreur LM Studio : {e}")
        return []

# ─── Téléchargement des PDFs ──────────────────────────────────────────────────

def download_pdfs(markdown_text, page_url, save_dir):
    """Trouve et télécharge les PDFs liés sur la page."""
    save_dir.mkdir(exist_ok=True)
    base = f"{urlparse(page_url).scheme}://{urlparse(page_url).netloc}"

    # Chercher toutes les URLs PDF dans le texte
    pdf_urls = set()
    pdf_urls.update(re.findall(r'https?://[^\s\)\]"\'<>]+\.pdf', markdown_text, re.IGNORECASE))
    for rel in re.findall(r'\]\(([^)]+\.pdf)\)', markdown_text, re.IGNORECASE):
        pdf_urls.add(urljoin(base, rel) if not rel.startswith("http") else rel)

    downloaded = []
    headers = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"}

    for pdf_url in list(pdf_urls)[:10]:  # Max 10 PDFs par page
        try:
            filename = re.sub(r"[^\w\-_.]", "_", pdf_url.split("/")[-1].split("?")[0])
            if not filename.lower().endswith(".pdf"):
                filename += ".pdf"
            filepath = save_dir / filename

            if not filepath.exists():
                r = requests.get(pdf_url, timeout=15, headers=headers, stream=True)
                if r.status_code == 200 and "pdf" in r.headers.get("content-type", "").lower():
                    filepath.write_bytes(r.content)
                    print(f"    📄 PDF : {filename}")
            downloaded.append(str(filepath))
        except Exception:
            pass

    return downloaded

# ─── Scraping d'une URL ───────────────────────────────────────────────────────

class _HTMLTextExtractor(HTMLParser):
    """Extrait le texte brut d'une page HTML en ignorant scripts et styles."""
    def __init__(self):
        super().__init__()
        self._parts = []
        self._skip = False

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "nav", "footer", "head"):
            self._skip = True

    def handle_endtag(self, tag):
        if tag in ("script", "style", "nav", "footer", "head"):
            self._skip = False

    def handle_data(self, data):
        if not self._skip:
            text = data.strip()
            if text:
                self._parts.append(text)

    def get_text(self):
        return "\n".join(self._parts)


_REQUESTS_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "fr-FR,fr;q=0.9,en;q=0.5",
}


def scrape_with_requests(url):
    """Fallback HTTP simple quand crawl4ai échoue (sites anti-bot légers)."""
    try:
        r = requests.get(url, headers=_REQUESTS_HEADERS, timeout=20, allow_redirects=True)
        if r.status_code != 200:
            return None
        extractor = _HTMLTextExtractor()
        extractor.feed(r.text)
        text = re.sub(r"\s{3,}", "\n\n", extractor.get_text()).strip()
        return text if len(text) > 100 else None
    except Exception:
        return None


async def scrape_url(url, crawler):
    """Scrape une URL — tente crawl4ai d'abord, puis requests en fallback."""
    try:
        result = await crawler.arun(
            url=url,
            word_count_threshold=30,
            remove_overlay_elements=True,
            bypass_cache=True,
        )
        if result.success and result.markdown:
            return result.markdown
    except Exception as e:
        print(f"  ⚠️  Erreur crawl4ai : {e}")

    # Fallback : requête HTTP classique
    text = scrape_with_requests(url)
    if text:
        print(f"  ↩  Fallback HTTP utilisé")
    return text

# ─── Détection de sous-pages AAP ─────────────────────────────────────────────

# Mots-clés présents dans les liens vers des AAP individuels
_AAP_LINK_KEYWORDS = [
    "appel", "projet", "candidature", "candidater", "aap",
    "financement", "subvention", "dotation", "programme", "soutien",
    "appels", "projets",
]

# Mots à exclure (navigation générale, pas des AAP)
_EXCLUDE_PATTERNS = [
    "#", "mailto:", "javascript:", "login", "contact", "actualite",
    "presse", "newsletter", "rss", "sitemap", "cookie", "mentions",
    "faq", "aide", "accueil", "home",
]


def extract_subpage_links(markdown_text, base_url):
    """
    Extrait depuis une page de liste les liens qui pointent vers
    des sous-pages d'appels à projets individuels (même domaine).
    """
    parsed   = urlparse(base_url)
    domain   = parsed.netloc
    base_abs = f"{parsed.scheme}://{domain}"

    candidates = {}  # url → texte du lien

    # Liens absolus dans le markdown : [texte](https://...)
    for text, url in re.findall(r'\[([^\]]{3,120})\]\((https?://[^\)\s]{5,})\)', markdown_text):
        candidates[url.rstrip(")")] = text

    # Liens relatifs dans le markdown : [texte](/chemin...)
    for text, path in re.findall(r'\[([^\]]{3,120})\]\((/[^\)\s]{2,})\)', markdown_text):
        candidates[base_abs + path.rstrip(")")] = text

    aap_links = []
    for url, text in candidates.items():
        # Rester sur le même domaine
        if urlparse(url).netloc != domain:
            continue
        # Exclure les URLs de navigation générale
        url_lower = url.lower()
        if any(excl in url_lower for excl in _EXCLUDE_PATTERNS):
            continue
        # Garder si le texte du lien OU l'URL contient un mot-clé AAP
        combined = (text + " " + url).lower()
        if any(kw in combined for kw in _AAP_LINK_KEYWORDS):
            aap_links.append(url)

    # Dédupliquer en conservant l'ordre
    seen = set()
    result = []
    for url in aap_links:
        if url not in seen and url != base_url:
            seen.add(url)
            result.append(url)

    return result


# ─── Export Excel ─────────────────────────────────────────────────────────────

def export_to_excel():
    """Exporte toute la base SQLite vers un fichier Excel formaté."""
    conn = sqlite3.connect(DB_PATH)
    rows = conn.execute("""
        SELECT titre, description, date_debut, date_fin,
               criteres_eligibilite, thematique, url_source, date_ajout
        FROM aap
        ORDER BY date_ajout DESC
    """).fetchall()
    total = conn.execute("SELECT COUNT(*) FROM aap").fetchone()[0]
    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Appels à Projets"

    headers = [
        "Titre", "Description", "Date début", "Date fin",
        "Critères d'éligibilité", "Thématique", "URL source", "Date d'ajout"
    ]

    # En-têtes colorés
    fill   = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    font   = Font(color="FFFFFF", bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = fill
        cell.font = font
        cell.alignment = center

    ws.row_dimensions[1].height = 25

    # Données
    wrap = Alignment(vertical="top", wrap_text=True)
    for row_data in rows:
        ws.append(row_data)
        ws.row_dimensions[ws.max_row].height = 40
        for cell in ws[ws.max_row]:
            cell.alignment = wrap

    # Largeurs des colonnes
    col_widths = [50, 60, 14, 14, 50, 25, 50, 18]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = width

    wb.save(EXCEL_PATH)
    print(f"  → Excel exporté : {EXCEL_PATH} ({total} entrées au total)")

# ─── Programme principal ──────────────────────────────────────────────────────

async def main():
    now = datetime.now()
    print(f"\n{'='*60}")
    print(f"  Scraping AAP Santé — {now.strftime('%d/%m/%Y %H:%M')}")
    print(f"{'='*60}\n")

    # Initialisation
    PDFS_DIR.mkdir(exist_ok=True)
    init_db()

    urls = load_urls(URLS_FILE)
    if not urls:
        print("Aucune URL à scraper. Vérifiez le fichier urls.txt")
        return

    print(f"📋 {len(urls)} URLs chargées depuis {URLS_FILE}")
    print(f"🤖 LM Studio : {LM_STUDIO_URL}")
    print(f"💾 Base de données : {DB_PATH}\n")

    total_new    = 0
    total_found  = 0
    total_errors = 0

    async with AsyncWebCrawler(verbose=False) as crawler:
        for i, url in enumerate(urls, 1):
            print(f"[{i:02d}/{len(urls)}] {url}")

            # ── Étape 1 : scraper la page de liste ──────────────────────────
            markdown = await scrape_url(url, crawler)
            if not markdown:
                print(f"  ✗ Page inaccessible\n")
                total_errors += 1
                continue

            # PDFs de la page principale
            fichiers_liste = download_pdfs(markdown, url, PDFS_DIR)

            # ── Étape 2 : détecter les liens vers des sous-pages AAP ────────
            subpage_links = extract_subpage_links(markdown, url)

            if subpage_links:
                print(f"  → {len(subpage_links)} sous-page(s) AAP détectée(s), exploration...")
                page_found = 0
                page_new   = 0

                for sub_url in subpage_links[:25]:  # max 25 sous-pages par site
                    sub_markdown = await scrape_url(sub_url, crawler)
                    if not sub_markdown:
                        continue

                    sub_fichiers = download_pdfs(sub_markdown, sub_url, PDFS_DIR)
                    tous_fichiers = fichiers_liste + sub_fichiers

                    sub_aap = extract_with_lm_studio(sub_markdown, sub_url)
                    if sub_aap:
                        new = save_to_db(sub_aap, sub_url, tous_fichiers)
                        page_found += len(sub_aap)
                        page_new   += new
                        print(f"    ✓ {sub_url.split('/')[-1]} : {len(sub_aap)} AAP ({new} nouveaux)")

                total_found += page_found
                total_new   += page_new
                if page_found == 0:
                    print(f"  → Aucun appel à projet détecté dans les sous-pages")
                print()

            else:
                # Pas de sous-pages trouvées → extraire directement la page
                print(f"  → Extraction directe en cours...")
                aap_list = extract_with_lm_studio(markdown, url)

                if not aap_list:
                    print(f"  → Aucun appel à projet détecté\n")
                    continue

                fichiers_liste += download_pdfs(markdown, url, PDFS_DIR)
                new = save_to_db(aap_list, url, fichiers_liste)
                total_found += len(aap_list)
                total_new   += new
                print(f"  ✓ {len(aap_list)} AAP trouvés — {new} nouveaux ajoutés\n")

    # 5. Exporter en Excel
    print("Export Excel en cours...")
    export_to_excel()

    # Résumé final
    print(f"\n{'='*60}")
    print(f"  ✅ Scraping terminé — {now.strftime('%d/%m/%Y %H:%M')}")
    print(f"  → {total_found} appels trouvés au total")
    print(f"  → {total_new} nouveaux ajoutés à la base")
    print(f"  → {total_errors} site(s) inaccessible(s)")
    print(f"  → PDFs dans le dossier : {PDFS_DIR}/")
    print(f"  → Base de données      : {DB_PATH}")
    print(f"  → Fichier Excel        : {EXCEL_PATH}")
    print(f"{'='*60}\n")


if __name__ == "__main__":
    asyncio.run(main())
